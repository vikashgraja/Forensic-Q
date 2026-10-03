"""
Q-Bank Forensic Unit & Integration Tests
Tests statement parsing, entity extraction, cash deposit detection, Hyundai rules, and Tabulator APIs.
"""

import io
import uuid
from datetime import UTC, datetime
from decimal import Decimal

import pandas as pd
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from docx import Document

from .backend.statement_parser import (
    extract_clean_tracking_name,
    extract_tables_from_word,
    format_inr,
    normalize_dataframe,
    parse_bank_statement_dataframe,
)
from .models import AuditedPerson, BankAccount, BankTransaction, WatchlistRule
from .selectors import (
    fuzzy_search_transactions,
    get_all_statement_transactions,
    get_bank_dashboard_metrics,
    get_frequent_counterparties,
    get_frequent_transactions_breakdown,
    get_hyundai_metrics,
)
from .services import (
    delete_audited_person,
    delete_bank_account,
    ingest_bank_statement_file,
)


class QBankStatementTests(TestCase):
    """
    Validates parsing, normalization, and entity extraction.
    """

    def test_format_inr(self):
        self.assertEqual(format_inr(1000), "1,000.00")
        self.assertEqual(format_inr(150000), "1,50,000.00")
        self.assertEqual(format_inr(12345678.50), "1,23,45,678.50")
        self.assertEqual(format_inr(-50000), "-50,000.00")

    def test_extract_clean_tracking_name(self):
        # UPI VPA
        name1 = extract_clean_tracking_name("UPI/428912/john.doe@okaxis/Payment for supplies")
        self.assertEqual(name1.upper(), "JOHN DOE")

        # Slash-delimited NetBanking
        name2 = extract_clean_tracking_name("NEFT/N1029384/ABC LOGISTICS CORP/HDFC000123")
        self.assertEqual(name2.upper(), "ABC LOGISTICS CORP")

        # POS / Card merchant
        name3 = extract_clean_tracking_name("POS-RESTAURANT CHENNAI 12:30:00")
        self.assertIn("RESTAURANT", name3.upper())

    def test_normalize_dataframe(self):
        raw_data = {
            "Tran Date": ["2026-03-01", "2026-03-02"],
            "Particulars": ["UPI/test@okaxis", "NEFT/SUPPLIER/123"],
            "Withdrawal (Dr)": ["1,500.00", "0.00"],
            "Deposit (Cr)": ["0.00", "25,000.00"],
            "Balance": ["1,00,000.00", "1,25,000.00"],
        }
        df = pd.DataFrame(raw_data)
        normalized = normalize_dataframe(df)

        self.assertIn("Date", normalized.columns)
        self.assertIn("Narration", normalized.columns)
        self.assertIn("Debit Amount", normalized.columns)
        self.assertIn("Credit Amount", normalized.columns)
        self.assertIn("Closing Balance", normalized.columns)
        self.assertEqual(normalized["Debit Amount"].iloc[0], 1500.00)
        self.assertEqual(normalized["Credit Amount"].iloc[1], 25000.00)

    def test_extract_clean_tracking_name_edge_cases(self):
        # Numeric VPA left side
        name = extract_clean_tracking_name("UPI/12345/9988776655@paytm/Supermarket Supplies")
        self.assertIn("SUPERMARKET", name.upper())

        # POS with time and state code
        name_pos = extract_clean_tracking_name("POS-PURCHASE CHENNAI TN 14:22:10 RESTAURANT")
        self.assertTrue(len(name_pos) > 0)

        # Fallback words
        name_words = extract_clean_tracking_name("MISC TRANSFER FROM CLIENT")
        self.assertIn("MISC", name_words.upper())

        # Fallback Other Account Operations
        name_fallback = extract_clean_tracking_name("123456 999999")
        self.assertEqual(name_fallback, "Other Account Operations")

    def test_format_inr_edge_cases(self):
        self.assertEqual(format_inr("not_a_number"), "0.00")
        self.assertEqual(format_inr(50), "50.00")
        self.assertEqual(format_inr(-150000), "-1,50,000.00")

    def test_extract_tables_from_word(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.cell(0, 0).text = "Date"
        table.cell(0, 1).text = "Narration"
        table.cell(0, 2).text = "Amount"
        table.cell(1, 0).text = "01/01/2026"
        table.cell(1, 1).text = "Consulting Fee"
        table.cell(1, 2).text = "50000"

        doc_buf = io.BytesIO()
        doc.save(doc_buf)
        doc_buf.seek(0)

        df = extract_tables_from_word(doc_buf)
        self.assertFalse(df.empty)

    def test_parse_bank_statement_dataframe_various_types(self):
        # Excel buffer
        df_in = pd.DataFrame(
            {
                "Date": ["01/01/2026"],
                "Narration": ["UPI/test@okaxis/Vendor"],
                "Debit Amount": [1000.0],
                "Credit Amount": [0.0],
                "Closing Balance": [50000.0],
                "page_num": [1],
            }
        )
        xl_buf = io.BytesIO()
        with pd.ExcelWriter(xl_buf, engine="openpyxl") as writer:
            df_in.to_excel(writer, index=False)
        xl_buf.seek(0)

        df_out = parse_bank_statement_dataframe(xl_buf, "statement.xlsx")
        self.assertIn("UPI_Name", df_out.columns)
        self.assertNotIn("page_num", df_out.columns)

        # Unknown extension fallback to CSV
        csv_buf = io.BytesIO(
            b"Date,Narration,Debit Amount,Credit Amount,Closing Balance\n01/01/2026,Cash Deposit,0,5000,55000\n"
        )
        df_csv = parse_bank_statement_dataframe(csv_buf, "unknown_data.dat")
        self.assertIn("Narration", df_csv.columns)


class QBankServicesAndSelectorsTests(TestCase):
    """
    Validates atomic ingestion, risk scoring, selectors, and API endpoints.
    """

    def setUp(self):
        self.client = Client()
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        # Create sample CSV statement in memory
        csv_content = """Date,Narration,Transaction ID,Debit Amount,Credit Amount,Closing Balance
01/01/2026,UPI/998811/sarla_enterprises@okaxis/Invoice 101,TXN1001,500000.00,0.00,1000000.00
02/01/2026,BY CASH DEPOSIT CDM BRANCH,TXN1002,0.00,75000.00,1075000.00
03/01/2026,NEFT/HMIL VENDOR PAYMENT HYUNDAI MOTORS,TXN1003,0.00,250000.00,1325000.00
04/01/2026,UPI/998811/sarla_enterprises@okaxis/Invoice 102,TXN1004,15000.00,0.00,1310000.00
05/01/2026,UPI/998811/sarla_enterprises@okaxis/Invoice 103,TXN1005,20000.00,0.00,1290000.00
"""
        csv_file = io.BytesIO(csv_content.encode("utf-8"))

        WatchlistRule.objects.create(rule_name="Sarla Flag", keyword="sarla", risk_weight=40)

        self.account = ingest_bank_statement_file(
            file_obj_or_path=csv_file,
            filename="hdfc_statement.csv",
            account_holder="Vikash Test Account",
            bank_name="HDFC Bank",
            statement_label="Q1 Audit Investigation",
            account_number="987654321",
        )

    def test_ingest_bank_statement_metrics(self):
        self.assertEqual(self.account.total_transactions, 5)
        self.assertEqual(self.account.cash_deposit_count, 1)
        self.assertEqual(self.account.hyundai_count, 1)
        self.assertGreater(self.account.total_debit, Decimal("500000.00"))

        # Verify transaction risk scoring
        txns = BankTransaction.objects.filter(account=self.account)
        self.assertEqual(txns.count(), 5)

        # CDM cash deposit check
        cdm_txn = txns.filter(is_cash_deposit=True).first()
        self.assertIsNotNone(cdm_txn)
        self.assertEqual(cdm_txn.credit_amount, Decimal("75000.00"))

        # Hyundai check
        hyundai_txn = txns.filter(is_hyundai_related=True).first()
        self.assertIsNotNone(hyundai_txn)

    def test_selectors(self):
        metrics = get_bank_dashboard_metrics()
        self.assertEqual(metrics["total_accounts"], 1)
        self.assertEqual(metrics["total_txns"], 5)
        self.assertEqual(metrics["cash_deposit_count"], 1)
        self.assertEqual(metrics["hyundai_count"], 1)

        frequent = get_frequent_counterparties(account_id=self.account.id, min_interactions=2)
        self.assertTrue(len(frequent) > 0)
        self.assertIn("sarla", frequent[0]["party_name"].lower())

        hyundai_m = get_hyundai_metrics(account_id=self.account.id)
        self.assertTrue(hyundai_m["hyundai_present"])
        self.assertEqual(hyundai_m["total_count"], 1)

        # CDM Rows test
        from .selectors import (
            get_cdm_transactions,
            get_frequent_transactions_breakdown,
            get_hyundai_details,
        )

        cdm_rows = get_cdm_transactions(account_id=self.account.id)
        self.assertEqual(len(cdm_rows), 1)
        self.assertEqual(cdm_rows[0]["credit_amount"], 75000.0)

        # Hyundai details test
        h_details = get_hyundai_details(account_id=self.account.id)
        self.assertTrue(h_details["hyundai_present"])
        self.assertEqual(len(h_details["rows"]), 1)

        # Frequent breakdown test
        breakdown = get_frequent_transactions_breakdown(
            account_id=self.account.id, min_transactions=2
        )
        self.assertIn("debit_choices", breakdown)
        self.assertIn("credit_choices", breakdown)
        self.assertTrue(len(breakdown["debit_choices"]) > 0)

    def test_fuzzy_search(self):
        matches = fuzzy_search_transactions(
            account_id=self.account.id,
            keywords_str="sarla, trust",
            threshold=70,
        )
        self.assertGreater(len(matches), 0)
        self.assertIn("sarla", matches[0]["narration"].lower())

    def test_api_transactions(self):
        res = self.client.get(
            reverse("q_bank:transactions_api"), {"account_id": str(self.account.id)}
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["total_count"], 5)
        self.assertEqual(len(data["data"]), 5)

    def test_views_and_exports(self):
        dash_res = self.client.get(reverse("q_bank:dashboard"))
        self.assertEqual(dash_res.status_code, 200)
        self.assertContains(dash_res, "Q-Bank")

        detail_res = self.client.get(
            reverse("q_bank:account_detail", args=[self.account.id]), follow=True
        )
        self.assertEqual(detail_res.status_code, 200)
        self.assertContains(detail_res, "Vikash Test Account")

        export_res = self.client.get(reverse("q_bank:export_ledger_excel"))
        self.assertEqual(export_res.status_code, 200)
        self.assertEqual(
            export_res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_multi_bank_statements_per_person(self):
        person = self.account.person
        self.assertIsNotNone(person)
        self.assertEqual(person.full_name, "Vikash Test Account")

        # Ingest a second statement (SBI) for the same person
        sbi_csv = """Date,Narration,Transaction ID,Debit Amount,Credit Amount,Closing Balance
10/01/2026,UPI/112233/vendor_corp@sbi/Equip 201,SBI001,12000.00,0.00,50000.00
11/01/2026,BY CASH DEPOSIT CDM SBI KOYAMBEDU,SBI002,0.00,30000.00,80000.00
"""
        sbi_file = io.BytesIO(sbi_csv.encode("utf-8"))
        sbi_acc = ingest_bank_statement_file(
            file_obj_or_path=sbi_file,
            filename="sbi_statement.csv",
            person_id=person.id,
            bank_name="State Bank of India",
            statement_label="SBI Current Account",
            account_number="330011223344",
        )
        self.assertEqual(sbi_acc.person_id, person.id)
        self.assertEqual(person.bank_accounts.count(), 2)

        # Test Person detail view (Combined All Statements)
        person_res = self.client.get(reverse("q_bank:person_detail", args=[person.id]))
        self.assertEqual(person_res.status_code, 200)
        self.assertContains(person_res, "Vikash Test Account")
        self.assertContains(person_res, "HDFC Bank")
        self.assertContains(person_res, "State Bank of India")

        # Test Person detail view scoped to single account
        scoped_res = self.client.get(
            f"{reverse('q_bank:person_detail', args=[person.id])}?account_id={sbi_acc.id}"
        )
        self.assertEqual(scoped_res.status_code, 200)

        # Test API with person_id
        api_res = self.client.get(reverse("q_bank:transactions_api"), {"person_id": str(person.id)})
        self.assertEqual(api_res.status_code, 200)
        self.assertEqual(api_res.json()["total_count"], 7)  # 5 from HDFC + 2 from SBI

    def test_delete_person(self):
        person = self.account.person
        self.assertIsNotNone(person)

        success = delete_audited_person(person.id)
        self.assertTrue(success)
        self.assertEqual(BankAccount.objects.count(), 0)
        self.assertEqual(BankTransaction.objects.count(), 0)

    def test_delete_account(self):
        success = delete_bank_account(self.account.id)
        self.assertTrue(success)
        self.assertEqual(BankAccount.objects.count(), 0)
        self.assertEqual(BankTransaction.objects.count(), 0)

    def test_upload_statement_view_csv(self):
        csv_data = b"Date,Narration,Transaction ID,Debit Amount,Credit Amount,Closing Balance\n01/02/2026,UPI/8899/vendor@axis/Parts,TXN5501,12000,0,88000\n"
        upload_file = SimpleUploadedFile("axis_statement.csv", csv_data, content_type="text/csv")
        res = self.client.post(
            reverse("q_bank:upload_statement"),
            data={
                "statement_file": upload_file,
                "account_holder": "Rajesh Kumar",
                "bank_name": "Axis Bank",
                "statement_label": "Vendor Ingestion",
                "account_number": "9182736450",
            },
        )
        self.assertEqual(res.status_code, 302)
        acc = BankAccount.objects.filter(account_holder="Rajesh Kumar").first()
        self.assertIsNotNone(acc)
        self.assertEqual(acc.total_transactions, 1)

    def test_upload_statement_view_excel(self):
        df = pd.DataFrame(
            {
                "Date": ["01/02/2026"],
                "Narration": ["NEFT/SUPPLIER PAYMENT"],
                "Transaction ID": ["EXCEL001"],
                "Debit Amount": [45000.0],
                "Credit Amount": [0.0],
                "Closing Balance": [155000.0],
            }
        )
        xl_buf = io.BytesIO()
        with pd.ExcelWriter(xl_buf, engine="openpyxl") as writer:
            df.to_excel(writer, index=False)
        xl_buf.seek(0)
        upload_file = SimpleUploadedFile(
            "icici_statement.xlsx",
            xl_buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        res = self.client.post(
            reverse("q_bank:upload_statement"),
            data={
                "statement_file": upload_file,
                "account_holder": "ICICI Corporate",
                "bank_name": "ICICI Bank",
                "statement_label": "Excel Import",
            },
        )
        self.assertEqual(res.status_code, 302)
        acc = BankAccount.objects.filter(account_holder="ICICI Corporate").first()
        self.assertIsNotNone(acc)
        self.assertEqual(acc.total_transactions, 1)

    def test_upload_statement_auto_extracts_period_from_table(self):
        csv_data = (
            b"Date,Narration,Transaction ID,Debit Amount,Credit Amount,Closing Balance\n"
            b"01/03/2026,UPI/Vendor1,TXN101,5000,0,95000\n"
            b"28/03/2026,SALARY/Credit,TXN102,0,50000,145000\n"
        )
        upload_file = SimpleUploadedFile("hdfc_period_test.csv", csv_data, content_type="text/csv")
        res = self.client.post(
            reverse("q_bank:upload_statement"),
            data={
                "statement_file": upload_file,
                "account_holder": "Period Test Person",
                "bank_name": "HDFC Bank",
                "account_number": "501002984123",
            },
        )
        self.assertEqual(res.status_code, 302)
        acc = BankAccount.objects.filter(account_holder="Period Test Person").first()
        self.assertIsNotNone(acc)
        self.assertEqual(acc.total_transactions, 2)
        self.assertEqual(acc.statement_label, "01 Mar 2026 - 28 Mar 2026")

    def test_upload_statement_view_errors(self):
        # Missing file
        res_no_file = self.client.post(reverse("q_bank:upload_statement"), data={})
        self.assertEqual(res_no_file.status_code, 302)

        # Corrupt file
        bad_file = SimpleUploadedFile("corrupt.csv", b"\xff\xfe\x00\x00", content_type="text/csv")
        res_bad = self.client.post(
            reverse("q_bank:upload_statement"),
            data={"statement_file": bad_file, "account_holder": "Bad Data"},
        )
        self.assertEqual(res_bad.status_code, 302)

    def test_create_and_delete_person_views(self):
        # Create person view success
        create_res = self.client.post(
            reverse("q_bank:create_person"),
            data={
                "full_name": "Devi Prasad",
                "employee_id": "EMP-990",
                "department": "Finance",
                "designation": "Manager",
                "pan_number": "ABCDE1234F",
                "email": "devi.prasad@company.com",
            },
        )
        self.assertEqual(create_res.status_code, 302)
        person = AuditedPerson.objects.filter(full_name="Devi Prasad").first()
        self.assertIsNotNone(person)
        self.assertEqual(person.employee_id, "EMP-990")

        # Create person view with empty full name
        fail_res = self.client.post(reverse("q_bank:create_person"), data={"full_name": ""})
        self.assertEqual(fail_res.status_code, 302)

        # Delete person view success
        del_res = self.client.post(reverse("q_bank:delete_person", args=[person.id]))
        self.assertEqual(del_res.status_code, 302)
        self.assertFalse(AuditedPerson.objects.filter(id=person.id).exists())

        # Delete non-existent person
        bad_del = self.client.post(reverse("q_bank:delete_person", args=[uuid.uuid4()]))
        self.assertEqual(bad_del.status_code, 302)

    def test_account_detail_view_auto_link_and_404(self):
        # Create unlinked account
        unlinked_acc = BankAccount.objects.create(
            account_holder="Autonomous Custodian",
            bank_name="Canara Bank",
            account_number="CN102938",
            person=None,
        )
        res = self.client.get(reverse("q_bank:account_detail", args=[unlinked_acc.id]))
        self.assertEqual(res.status_code, 302)
        unlinked_acc.refresh_from_db()
        self.assertIsNotNone(unlinked_acc.person)
        self.assertEqual(unlinked_acc.person.full_name, "Autonomous Custodian")

        # 404 for non-existent account
        res_404 = self.client.get(reverse("q_bank:account_detail", args=[uuid.uuid4()]))
        self.assertEqual(res_404.status_code, 404)

    def test_delete_account_view_redirects(self):
        target_acc_id = self.account.id
        person_id = self.account.person.id
        del_res = self.client.post(reverse("q_bank:delete_account", args=[target_acc_id]))
        self.assertEqual(del_res.status_code, 302)
        self.assertIn(f"/bank/person/{person_id}/", del_res.url)

        # Delete non-existent account
        bad_del = self.client.post(reverse("q_bank:delete_account", args=[uuid.uuid4()]))
        self.assertEqual(bad_del.status_code, 302)

    def test_fuzzy_search_api_and_export_frequent(self):
        # Fuzzy search API with keywords
        fuzzy_url = reverse("q_bank:fuzzy_search_api")
        res = self.client.get(
            f"{fuzzy_url}?account_id={self.account.id}&keywords=sarla&threshold=50"
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "ok")
        self.assertGreaterEqual(data["total_matches"], 1)

        # Fuzzy search with person_id and edge case empty keywords
        self.assertEqual(
            fuzzy_search_transactions(person_id=self.account.person.id, keywords_str=""), []
        )
        matches_person = fuzzy_search_transactions(
            person_id=self.account.person.id, keywords_str="hyundai, , !!", threshold=60
        )
        self.assertIsInstance(matches_person, list)

        # Frequent transactions breakdown with multiple credit transactions
        BankTransaction.objects.create(
            account=self.account,
            txn_date=datetime.now(UTC),
            value_date=datetime.now(UTC),
            narration="CREDIT PAYMENT VENDOR ALPHA",
            party_name="Vendor Alpha",
            credit_amount=Decimal("50000.00"),
            closing_balance=Decimal("200000.00"),
        )
        BankTransaction.objects.create(
            account=self.account,
            txn_date=datetime.now(UTC),
            value_date=datetime.now(UTC),
            narration="SECOND CREDIT PAYMENT VENDOR ALPHA",
            party_name="Vendor Alpha",
            credit_amount=Decimal("75000.00"),
            closing_balance=Decimal("275000.00"),
        )
        breakdown = get_frequent_transactions_breakdown(
            account_id=self.account.id, min_transactions=2
        )
        self.assertIn("credit_groups", breakdown)
        self.assertIn("Vendor Alpha", breakdown["credit_groups"])

        # Export frequent counterparties Excel
        export_url = reverse("q_bank:export_frequent_excel")
        exp_res = self.client.get(f"{export_url}?account_id={self.account.id}")
        self.assertEqual(exp_res.status_code, 200)
        self.assertEqual(
            exp_res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_get_all_statement_transactions(self):
        person = self.account.person
        # 1. Filter by account_id
        txns_acc = get_all_statement_transactions(account_id=self.account.id)
        self.assertEqual(len(txns_acc), 5)
        sample = txns_acc[0]
        self.assertIn("id", sample)
        self.assertIn("date", sample)
        self.assertIn("narration", sample)
        self.assertIn("closing_balance_formatted", sample)
        self.assertIsInstance(sample["credit_amount"], float)
        self.assertIsInstance(sample["debit_amount"], float)

        # 2. Filter by person_id
        txns_person = get_all_statement_transactions(person_id=person.id)
        self.assertEqual(len(txns_person), 5)

        # 3. Limit parameter
        txns_limited = get_all_statement_transactions(person_id=person.id, limit=2)
        self.assertEqual(len(txns_limited), 2)

        # 4. Non-existent filter returns empty list
        empty_res = get_all_statement_transactions(account_id=uuid.uuid4())
        self.assertEqual(len(empty_res), 0)

    def test_person_detail_view_closing_balance_and_all_txns(self):
        person = self.account.person
        # 1. View without specific account (aggregate across accounts)
        url = reverse("q_bank:person_detail", kwargs={"person_id": person.id})
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertIn("all_transactions", res.context)
        self.assertGreaterEqual(len(res.context["all_transactions"]), 5)
        self.assertIn("closing_balance", res.context["view_metrics"])
        self.assertIn("closing_balance_formatted", res.context["view_metrics"])
        self.assertEqual(
            res.context["view_metrics"]["closing_balance"],
            Decimal("1290000.00"),
        )

        # 2. View with selected account_id query param
        res_acc = self.client.get(f"{url}?account_id={self.account.id}")
        self.assertEqual(res_acc.status_code, 200)
        self.assertEqual(res_acc.context["selected_account"].id, self.account.id)
        self.assertEqual(
            res_acc.context["view_metrics"]["closing_balance"],
            Decimal("1290000.00"),
        )

        # 3. Person with an account that has no transactions
        empty_acc = BankAccount.objects.create(
            person=person,
            bank_name="Empty Bank",
            account_number="ACC-EMPTY-001",
        )
        res_empty = self.client.get(f"{url}?account_id={empty_acc.id}")
        self.assertEqual(res_empty.status_code, 200)
        self.assertEqual(res_empty.context["view_metrics"]["closing_balance"], Decimal("0.00"))

    def test_parse_keywords_api_and_file_fuzzy_search(self):
        # 1. Test parse_keywords_api with plain text file
        txt_content = b"# Forensic Watchlist\ntrust\nsarla\nhawala\n# Comments\n"
        txt_file = SimpleUploadedFile("keywords.txt", txt_content, content_type="text/plain")
        parse_url = reverse("q_bank:parse_keywords_api")
        res = self.client.post(parse_url, {"file": txt_file})
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "ok")
        self.assertEqual(data["filename"], "keywords.txt")
        self.assertEqual(data["total_count"], 3)
        self.assertIn("trust", data["keywords"])
        self.assertIn("sarla", data["keywords"])
        self.assertIn("hawala", data["keywords"])

        # 2. Test parse_keywords_api with Excel file
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Keywords"
        ws.append(["Keyword", "Category"])
        ws.append(["sarla", "Entity"])
        ws.append(["trust", "Entity"])
        ws.append(["kickback", "Risk"])
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        excel_file = SimpleUploadedFile(
            "test_terms.xlsx", bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        res_xl = self.client.post(parse_url, {"file": excel_file})
        self.assertEqual(res_xl.status_code, 200)
        data_xl = res_xl.json()
        self.assertEqual(data_xl["status"], "ok")
        self.assertIn("sarla", data_xl["keywords"])
        self.assertIn("trust", data_xl["keywords"])
        self.assertIn("kickback", data_xl["keywords"])

        # 3. Test fuzzy_search_api with uploaded file
        fuzzy_url = reverse("q_bank:fuzzy_search_api")
        txt_search_file = SimpleUploadedFile("search.txt", b"sarla\ntrust", content_type="text/plain")
        res_fuzzy_file = self.client.post(
            fuzzy_url,
            {"account_id": str(self.account.id), "file": txt_search_file, "threshold": 50},
        )
        self.assertEqual(res_fuzzy_file.status_code, 200)
        fuzzy_data = res_fuzzy_file.json()
        self.assertEqual(fuzzy_data["status"], "ok")
        self.assertGreaterEqual(fuzzy_data["total_matches"], 1)
        self.assertIn("matched_keyword", fuzzy_data["matches"][0])
        self.assertTrue(fuzzy_data["matches"][0]["matched_keyword"])

