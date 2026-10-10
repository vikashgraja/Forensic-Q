import hashlib
import io
import json
import logging
import tempfile
import uuid
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock, patch

from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse

from core.db import configure_database_connection
from core.file_uploader import FileUploader
from core.logging import InterceptHandler, setup_logging
from core.models import InvestigationProfile
from core.modules import get_discovered_modules


class CoreFileUploaderTests(TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.uploader = FileUploader(self.temp_dir.name, default_ext=".bin")

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_single_chunk_upload_and_sha256(self):
        data = b"FORENSIC_EVIDENCE_PAYLOAD_CHUNK_12345"
        expected_sha = hashlib.sha256(data).hexdigest()

        res = self.uploader.append_chunk(
            upload_id="test_upload_01",
            chunk_index=0,
            total_chunks=1,
            chunk_data=data,
        )

        self.assertTrue(res["is_completed"])
        self.assertEqual(res["file_sha256"], expected_sha)
        self.assertEqual(res["current_size_bytes"], len(data))
        self.assertTrue(Path(res["file_path"]).exists())

    def test_multi_chunk_assembly(self):
        chunk1 = b"PART_ONE_"
        chunk2 = b"PART_TWO_"
        chunk3 = b"PART_THREE"
        full_data = chunk1 + chunk2 + chunk3
        expected_sha = hashlib.sha256(full_data).hexdigest()

        res1 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=0, total_chunks=3, chunk_data=chunk1
        )
        self.assertFalse(res1["is_completed"])

        res2 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=1, total_chunks=3, chunk_data=chunk2
        )
        self.assertFalse(res2["is_completed"])

        res3 = self.uploader.append_chunk(
            upload_id="multi_01", chunk_index=2, total_chunks=3, chunk_data=chunk3
        )
        self.assertTrue(res3["is_completed"])
        self.assertEqual(res3["file_sha256"], expected_sha)
        self.assertEqual(res3["current_size_bytes"], len(full_data))

    def test_extension_normalization_and_chunk0_overwrite(self):
        res1 = self.uploader.append_chunk(
            upload_id="no_dot_ext",
            chunk_index=0,
            total_chunks=1,
            chunk_data=b"Initial",
            extension="dat",
        )
        self.assertTrue(res1["is_completed"])
        target = Path(res1["file_path"])
        self.assertTrue(target.exists())
        self.assertEqual(target.suffix, ".dat")

        res2 = self.uploader.append_chunk(
            upload_id="no_dot_ext",
            chunk_index=0,
            total_chunks=1,
            chunk_data=b"Overwritten",
            extension=".dat",
        )
        self.assertTrue(res2["is_completed"])
        self.assertEqual(target.read_bytes(), b"Overwritten")


class CorePortalAuthMiddlewareAndViewsTests(TestCase):
    def setUp(self):
        self.client = Client()

    def test_unauthenticated_redirect_to_login(self):
        res = self.client.get("/")
        self.assertEqual(res.status_code, 302)
        self.assertIn("/login/?next=/", res.url)

    def test_exempt_path_permits_anonymous_access(self):
        res = self.client.get("/login/")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Authorized Personnel Only", res.content)

    def test_login_failure_with_wrong_password(self):
        res = self.client.post("/login/", {"password": "wrong_password", "next": "/"})
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Invalid portal access key", res.content)
        self.assertFalse(self.client.session.get("portal_authenticated", False))

    def test_login_success_with_valid_password(self):
        valid_pwd = getattr(settings, "PORTAL_ACCESS_PASSWORD", "forensiq2026")
        res = self.client.post("/login/", {"password": valid_pwd, "next": "/bank/"})
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/bank/")
        self.assertTrue(self.client.session.get("portal_authenticated", False))

    def test_open_redirect_sanitization(self):
        valid_pwd = getattr(settings, "PORTAL_ACCESS_PASSWORD", "forensiq2026")
        # Attempt malicious off-site redirect
        res = self.client.post(
            "/login/", {"password": valid_pwd, "next": "https://malicious-site.com"}
        )
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/")

        res_double_slash = self.client.post(
            "/login/", {"password": valid_pwd, "next": "//malicious-site.com"}
        )
        self.assertEqual(res_double_slash.status_code, 302)
        self.assertEqual(res_double_slash.url, "/")

    def test_portal_logout_locks_workstation(self):
        # Authenticate first
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        # Logout
        res = self.client.get(reverse("portal_logout"))
        self.assertEqual(res.status_code, 302)
        self.assertEqual(res.url, "/login/")
        self.assertFalse(self.client.session.get("portal_authenticated", False))

    def test_landing_view_authenticated(self):
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        res = self.client.get("/")
        self.assertEqual(res.status_code, 200)
        self.assertIn(b"Investigation Platform", res.content)


class CoreModulesDiscoveryTests(TestCase):
    def test_get_discovered_modules_structure(self):
        modules = get_discovered_modules()
        self.assertIsInstance(modules, list)
        self.assertGreaterEqual(len(modules), 8)

        app_names = [m["app_name"] for m in modules]
        self.assertIn("q_bank", app_names)
        self.assertIn("q_mail", app_names)
        self.assertIn("q_voice", app_names)
        self.assertNotIn("q_timeline", app_names)

        # Check expected keys in each module
        for m in modules:
            self.assertIn("app_name", m)
            self.assertIn("name", m)
            self.assertIn("tag", m)
            self.assertIn("accent", m)
            self.assertIn("href", m)
            self.assertIn("tagline", m)

    def test_missing_apps_dir(self):
        with patch("core.modules.settings.BASE_DIR", Path("/non_existent_folder_xyz")):
            modules = get_discovered_modules()
            self.assertEqual(modules, [])

    def test_discovered_modules_fallback_attributes(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            apps_dir = temp_path / "apps"
            apps_dir.mkdir()

            # 1. Module without AppConfig or default spec
            custom_app = apps_dir / "q_custom_inspect"
            custom_app.mkdir()
            (custom_app / "__init__.py").write_text("", encoding="utf-8")

            # 2. Excluded module
            excluded_app = apps_dir / "q_timeline"
            excluded_app.mkdir()
            (excluded_app / "__init__.py").write_text("", encoding="utf-8")

            with patch("core.modules.settings.BASE_DIR", temp_path):
                modules = get_discovered_modules()
                self.assertEqual(len(modules), 1)
                mod = modules[0]
                self.assertEqual(mod["app_name"], "q_custom_inspect")
                self.assertEqual(mod["num"], "01")
                self.assertEqual(mod["name"], "Custom Inspect")


class CoreLoggingAndDatabaseTests(TestCase):
    def test_intercept_handler_emit(self):
        handler = InterceptHandler()
        # Normal log record
        record_info = logging.LogRecord(
            name="test_logger",
            level=logging.INFO,
            pathname=__file__,
            lineno=10,
            msg="Forensic audit test message",
            args=(),
            exc_info=None,
        )
        handler.emit(record_info)

        # Custom/unknown level record triggering ValueError in logger.level
        record_custom = logging.LogRecord(
            name="test_logger",
            level=99,
            pathname=__file__,
            lineno=20,
            msg="Custom level test message",
            args=(),
            exc_info=None,
        )
        record_custom.levelname = "CUSTOM_LEVEL_UNKNOWN_99"
        handler.emit(record_custom)

    def test_setup_logging_initialization(self):
        # Call setup_logging and verify it completes without error
        setup_logging()

    def test_configure_database_connection(self):
        # 1. Non-sqlite connection is skipped
        mock_pg_conn = MagicMock()
        mock_pg_conn.vendor = "postgresql"
        configure_database_connection(sender=None, connection=mock_pg_conn)
        mock_pg_conn.cursor.assert_not_called()

        # 2. SQLite connection executes pragmas
        mock_sqlite_conn = MagicMock()
        mock_sqlite_conn.vendor = "sqlite"
        mock_cursor = MagicMock()
        mock_sqlite_conn.cursor.return_value.__enter__.return_value = mock_cursor
        configure_database_connection(sender=None, connection=mock_sqlite_conn)
        self.assertGreaterEqual(mock_cursor.execute.call_count, 4)

        # 3. SQLite connection cursor exception handled gracefully
        mock_err_conn = MagicMock()
        mock_err_conn.vendor = "sqlite"
        mock_err_conn.cursor.side_effect = RuntimeError("Pragma lock error")
        configure_database_connection(sender=None, connection=mock_err_conn)


class CoreInvestigationProfilesTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.profile = InvestigationProfile.objects.create(
            full_name="Target Custodian A",
            employee_id="EMP-1001",
            department="Procurement",
            designation="Manager",
            email="custodian.a@example.com",
            phone="+91 9876543210",
            is_substantiated=True,
            status="ACTIVE",
            notes="Under observation",
            avatar_color="orange",
        )

    def test_get_all_profiles(self):
        from core.profiles import get_all_profiles

        profiles = get_all_profiles()
        self.assertGreaterEqual(len(profiles), 1)
        self.assertEqual(profiles.first().full_name, "Target Custodian A")

    def test_get_profile_by_id(self):
        from core.profiles import get_profile_by_id

        self.assertIsNone(get_profile_by_id(None))
        self.assertIsNone(get_profile_by_id("invalid-uuid"))
        found = get_profile_by_id(self.profile.id)
        self.assertIsNotNone(found)
        self.assertEqual(found.id, self.profile.id)

    def test_get_and_set_active_profile(self):
        from core.profiles import get_active_profile, set_active_profile

        request = self.factory.get("/")
        request.session = {}

        # Initially none
        self.assertIsNone(get_active_profile(request))

        # Set active
        active = set_active_profile(request, self.profile.id)
        self.assertIsNotNone(active)
        self.assertEqual(request.session.get("active_profile_id"), str(self.profile.id))
        self.assertEqual(get_active_profile(request).id, self.profile.id)

        # Clear active
        cleared = set_active_profile(request, None)
        self.assertIsNone(cleared)
        self.assertNotIn("active_profile_id", request.session)

    def test_create_investigation_profile(self):
        from core.profiles import create_investigation_profile

        with self.assertRaises(ValueError):
            create_investigation_profile(full_name="   ")

        p = create_investigation_profile(
            full_name="New Auditee B",
            employee_id="EMP-2002",
            department="Finance",
            is_substantiated=True,
            status="INVALID",  # invalid choice falls back to ACTIVE
        )
        self.assertEqual(p.full_name, "New Auditee B")
        self.assertTrue(p.is_substantiated)
        self.assertEqual(p.status, "ACTIVE")
        self.assertEqual(p.category, "EMPLOYEE")
        self.assertEqual(p.related_employee, "")

    def test_create_investigation_profile_categories(self):
        from core.profiles import create_investigation_profile

        # Vendor
        v = create_investigation_profile(
            full_name="Apex Logistics Inc",
            category="VENDOR",
            department="Transport",
        )
        self.assertEqual(v.category, "VENDOR")
        self.assertEqual(v.get_category_display(), "Vendor")
        d_v = v.to_dict()
        self.assertEqual(d_v["category"], "VENDOR")
        self.assertEqual(d_v["category_display"], "Vendor")

        # Relative of Employee with details
        rel = create_investigation_profile(
            full_name="Pooja Sharma",
            category="RELATIVE_OF_EMPLOYEE",
            related_employee="Spouse of Rajesh Sharma (EMP-102)",
        )
        self.assertEqual(rel.category, "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(rel.related_employee, "Spouse of Rajesh Sharma (EMP-102)")
        d_rel = rel.to_dict()
        self.assertEqual(d_rel["category"], "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(d_rel["related_employee"], "Spouse of Rajesh Sharma (EMP-102)")

        # Fallback invalid category to EMPLOYEE
        other = create_investigation_profile(
            full_name="Third Party Contractor",
            category="INVALID_TYPE",
        )
        self.assertEqual(other.category, "EMPLOYEE")

    def test_update_investigation_profile_categories(self):
        from core.profiles import update_investigation_profile

        updated = update_investigation_profile(
            self.profile.id,
            full_name="Target Custodian A",
            category="RELATIVE_OF_EMPLOYEE",
            related_employee="Brother of VP Procurement",
        )
        self.assertEqual(updated.category, "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(updated.related_employee, "Brother of VP Procurement")
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.category, "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(self.profile.related_employee, "Brother of VP Procurement")

    def test_resolve_or_create_profile_from_request(self):
        from core.profiles import resolve_or_create_profile_from_request, set_active_profile

        # 1. Existing Profile Selected
        req1 = self.factory.post("/", {"profile_id": str(self.profile.id)})
        req1.session = {}
        prof, name = resolve_or_create_profile_from_request(req1)
        self.assertEqual(prof.id, self.profile.id)
        self.assertEqual(name, self.profile.full_name)

        # 2. Inline New Profile
        req2 = self.factory.post(
            "/",
            {
                "new_profile_name": "Inline Person",
                "new_profile_dept": "IT",
                "new_profile_role": "Admin",
            },
        )
        req2.session = {}
        prof2, name2 = resolve_or_create_profile_from_request(req2)
        self.assertEqual(prof2.full_name, "Inline Person")
        self.assertEqual(name2, "Inline Person")

        # 2b. Inline Profile that already exists
        req2b = self.factory.post("/", {"new_profile_name": self.profile.full_name})
        req2b.session = {}
        prof2b, name2b = resolve_or_create_profile_from_request(req2b)
        self.assertEqual(prof2b.id, self.profile.id)

        # 3. Legacy Name Fallback
        req3 = self.factory.post("/", {"custodian_name": "Legacy Target"})
        req3.session = {}
        prof3, name3 = resolve_or_create_profile_from_request(req3, default_department="Logistics")
        self.assertEqual(prof3.full_name, "Legacy Target")
        self.assertEqual(prof3.department, "Logistics")

        # 3b. Legacy Name that already exists
        req3b = self.factory.post("/", {"account_holder": self.profile.full_name})
        req3b.session = {}
        prof3b, name3b = resolve_or_create_profile_from_request(req3b)
        self.assertEqual(prof3b.id, self.profile.id)

        # 4. Fallback to active session profile
        req4 = self.factory.post("/", {})
        req4.session = {}
        set_active_profile(req4, self.profile.id)
        prof4, name4 = resolve_or_create_profile_from_request(req4)
        self.assertEqual(prof4.id, self.profile.id)

        # 5. Empty
        req5 = self.factory.post("/", {})
        req5.session = {}
        prof5, name5 = resolve_or_create_profile_from_request(req5)
        self.assertIsNone(prof5)
        self.assertEqual(name5, "")

    def test_sync_all_existing_entities_to_profiles(self):
        from django.utils import timezone
        from q_bank.models import AuditedPerson
        from q_verify.models import VerificationCase
        from q_voice.models import AudioRecording

        from core.profiles import sync_all_existing_entities_to_profiles

        AuditedPerson.objects.create(full_name="Synced Bank Auditee", notes="flagged transaction")
        AudioRecording.objects.create(
            call_ref="CALL-SYNC-001",
            call_title="Synchronized Wiretap",
            call_timestamp=timezone.now(),
            custodian_name="Synced Voice Subject",
            risk_score=85,
        )
        VerificationCase.objects.create(
            custodian_name="Synced Verify Subject", custodian_department="HR"
        )

        created = sync_all_existing_entities_to_profiles()
        self.assertGreaterEqual(created, 3)


class CoreProfileViewsTests(TestCase):
    def setUp(self):
        self.client = Client()
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        self.profile = InvestigationProfile.objects.create(
            full_name="Target Person X",
            employee_id="EMP-9009",
            department="Operations",
        )

    def test_create_profile_view_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps({"full_name": "AJAX Created Person", "department": "Legal"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["profile"]["full_name"], "AJAX Created Person")

    def test_create_profile_view_json_empty_name(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps({"full_name": "  "}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)
        data = res.json()
        self.assertEqual(data["status"], "error")

    def test_create_profile_view_form(self):
        res = self.client.post(
            reverse("create_profile"),
            data={"full_name": "Form Created Person", "department": "Audit"},
        )
        self.assertEqual(res.status_code, 302)
        self.assertTrue(
            InvestigationProfile.objects.filter(full_name="Form Created Person").exists()
        )

    def test_create_profile_view_with_category_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps(
                {
                    "full_name": "Vendor Profile Alpha",
                    "category": "VENDOR",
                    "department": "Supply Chain",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["profile"]["category"], "VENDOR")
        self.assertEqual(data["profile"]["category_display"], "Vendor")

    def test_create_profile_view_with_relative_category_form(self):
        res = self.client.post(
            reverse("create_profile"),
            data={
                "full_name": "Kavita Verma",
                "category": "RELATIVE_OF_EMPLOYEE",
                "related_employee": "Daughter of GM Operations",
                "department": "External",
            },
        )
        self.assertEqual(res.status_code, 302)
        prof = InvestigationProfile.objects.filter(full_name="Kavita Verma").first()
        self.assertIsNotNone(prof)
        self.assertEqual(prof.category, "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(prof.related_employee, "Daughter of GM Operations")

    def test_set_active_profile_view_json(self):
        # Set active
        res1 = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": str(self.profile.id)}),
            content_type="application/json",
        )
        self.assertEqual(res1.status_code, 200)
        self.assertEqual(res1.json()["active_profile"]["id"], str(self.profile.id))

        # Clear active
        res2 = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": ""}),
            content_type="application/json",
        )
        self.assertEqual(res2.status_code, 200)
        self.assertIsNone(res2.json()["active_profile"])

    def test_create_profile_view_form_empty(self):
        res = self.client.post(
            reverse("create_profile"),
            data={"full_name": ""},
        )
        self.assertEqual(res.status_code, 302)

    def test_set_active_profile_view_form(self):
        res = self.client.post(
            reverse("set_active_profile"),
            data={"profile_id": str(self.profile.id)},
        )
        self.assertEqual(res.status_code, 302)

    def test_profile_list_api_view(self):
        res = self.client.get(reverse("api_profiles"))
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertGreaterEqual(len(data["profiles"]), 1)

    def test_create_profile_view_malformed_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=b"invalid-json{",
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)

    def test_set_active_profile_view_malformed_json(self):
        res = self.client.post(
            reverse("set_active_profile"),
            data=b"invalid-json{",
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)

    def test_set_active_profile_invalid_id(self):
        from core.profiles import set_active_profile

        request = RequestFactory().post("/")
        request.session = {}
        result = set_active_profile(request, "00000000-0000-0000-0000-000000000000")
        self.assertIsNone(result)

    def test_context_processor_error_fallback(self):
        from core.context_processors import global_profiles_context

        request = RequestFactory().get("/")
        request.session = {}
        with patch(
            "core.context_processors.get_all_profiles", side_effect=RuntimeError("DB offline")
        ):
            ctx = global_profiles_context(request)
            self.assertEqual(ctx["investigation_profiles"], [])
            self.assertEqual(ctx["total_profiles_count"], 0)

    def test_create_profile_with_keywords(self):
        from core.profiles import create_investigation_profile

        prof = create_investigation_profile(
            full_name="Keyword Subject",
            department="Risk & Audit",
            keywords=["Kickback", "commission", "KICKBACK", "off-book"],
        )
        self.assertEqual(prof.keywords, ["Kickback", "commission", "off-book"])
        self.assertIn("keywords", prof.to_dict())

    def test_add_keywords_to_profile(self):
        from core.profiles import add_keywords_to_profile, create_investigation_profile

        prof = create_investigation_profile(
            full_name="Investigation Target",
            keywords=["bribe"],
        )
        updated = add_keywords_to_profile(prof.id, "cash, gift, bribe, secret")
        self.assertEqual(updated.keywords, ["bribe", "cash", "gift", "secret"])

    def test_add_keywords_to_profile_not_found(self):
        from core.profiles import add_keywords_to_profile

        with self.assertRaises(ValueError):
            add_keywords_to_profile(uuid.uuid4(), "audit")

    def test_create_profile_view_with_keywords_json(self):
        res = self.client.post(
            reverse("create_profile"),
            data=json.dumps(
                {
                    "full_name": "Profile With Keywords",
                    "department": "Security",
                    "keywords": ["shell_company", "wire_transfer"],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["profile"]["keywords"], ["shell_company", "wire_transfer"])

    def test_create_profile_view_with_keywords_form(self):
        res = self.client.post(
            reverse("create_profile"),
            data={
                "full_name": "Form Profile Keywords",
                "department": "Procurement",
                "keywords": json.dumps(["cash_payment", "hawala"]),
            },
        )
        self.assertEqual(res.status_code, 302)
        prof = InvestigationProfile.objects.filter(full_name="Form Profile Keywords").first()
        self.assertIsNotNone(prof)
        self.assertEqual(prof.keywords, ["cash_payment", "hawala"])

    def test_add_profile_keywords_view_success(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": "fraud, evasion"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("fraud", data["profile"]["keywords"])
        self.assertIn("evasion", data["profile"]["keywords"])

    def test_add_profile_keywords_view_empty(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": ""}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)

    def test_add_profile_keywords_view_not_found(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": str(uuid.uuid4())})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": "audit"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 404)

    def test_parse_keywords_file_view_txt_success(self):
        txt_content = b'kickback, bribe\n"consulting fee"\toff-book;"secret commission"'
        uploaded_file = SimpleUploadedFile("terms.txt", txt_content, content_type="text/plain")
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("kickback", data["keywords"])
        self.assertIn("bribe", data["keywords"])
        self.assertIn("consulting fee", data["keywords"])
        self.assertIn("off-book", data["keywords"])
        self.assertIn("secret commission", data["keywords"])

    def test_parse_keywords_file_view_xlsx_rejected(self):
        uploaded_file = SimpleUploadedFile(
            "investigation_keywords.xlsx",
            b"dummy excel content",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        data = res.json()
        self.assertEqual(data["status"], "error")
        self.assertIn("Only .txt files are supported", data["message"])

    def test_parse_keywords_file_view_csv_rejected(self):
        uploaded_file = SimpleUploadedFile(
            "investigation_keywords.csv",
            b"keyword1,keyword2",
            content_type="text/csv",
        )
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        data = res.json()
        self.assertEqual(data["status"], "error")
        self.assertIn("Only .txt files are supported", data["message"])

    def test_parse_keywords_file_view_no_file(self):
        res = self.client.post(reverse("parse_keywords_file"), {})
        self.assertEqual(res.status_code, 400)
        self.assertIn("No file uploaded", res.json()["message"])

    def test_parse_keywords_file_view_invalid_extension(self):
        uploaded_file = SimpleUploadedFile(
            "malicious.exe", b"binary", content_type="application/octet-stream"
        )
        res = self.client.post(
            reverse("parse_keywords_file"),
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("Only .txt files are supported", res.json()["message"])

    def test_upload_profile_keywords_file_view_success(self):
        txt_content = b'"unauthorized payment", "phantom vendor"\t"off book deal"\r\nhawala'
        uploaded_file = SimpleUploadedFile("keywords.txt", txt_content, content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("unauthorized payment", data["profile"]["keywords"])
        self.assertIn("phantom vendor", data["profile"]["keywords"])
        self.assertIn("off book deal", data["profile"]["keywords"])
        self.assertIn("hawala", data["profile"]["keywords"])

        # Check DB updated
        self.profile.refresh_from_db()
        self.assertIn("unauthorized payment", self.profile.keywords)
        self.assertIn("off book deal", self.profile.keywords)

    def test_upload_profile_keywords_file_view_empty_file(self):
        uploaded_file = SimpleUploadedFile("empty.txt", b"   \n\n  ", content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("No valid keywords found", res.json()["message"])

    def test_upload_profile_keywords_file_view_not_found(self):
        uploaded_file = SimpleUploadedFile("test.txt", b"keyword1", content_type="text/plain")
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(uuid.uuid4())})
        res = self.client.post(
            url,
            {"file": uploaded_file},
        )
        self.assertEqual(res.status_code, 404)

    def test_upload_profile_keywords_csv_rejected(self):
        uploaded_file = SimpleUploadedFile(
            "test.csv", b"keyword1,keyword2", content_type="text/csv"
        )
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(url, {"file": uploaded_file})
        self.assertEqual(res.status_code, 400)
        self.assertIn("Only .txt files are supported", res.json()["message"])

    def test_extract_keywords_from_file_delimiters_and_quotes(self):
        from core.profiles import extract_keywords_from_file

        # Tests space, tab, comma, enter (\r\n and \n), semicolon delimiters and quotes
        txt_content = (
            b"kickback, bribe\t\"shell company\"\r\nhawala;off-book   'round tripping'  secret"
        )
        f = io.BytesIO(txt_content)
        keywords = extract_keywords_from_file(f, "keywords.txt")
        self.assertIn("kickback", keywords)
        self.assertIn("bribe", keywords)
        self.assertIn("shell company", keywords)
        self.assertIn("hawala", keywords)
        self.assertIn("off-book", keywords)
        self.assertIn("round tripping", keywords)
        self.assertIn("secret", keywords)

    def test_extract_keywords_from_file_non_utf8_txt(self):
        from core.profiles import extract_keywords_from_file

        # Test latin-1 encoded text file
        content = "bribe, fráud\tkickback\r\nsecret".encode("latin-1")
        f = io.BytesIO(content)
        keywords = extract_keywords_from_file(f, "latin.txt")
        self.assertIn("bribe", keywords)
        self.assertIn("fráud", keywords)
        self.assertIn("kickback", keywords)
        self.assertIn("secret", keywords)

    def test_edit_profile_view_json_success(self):
        url = reverse("edit_profile", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps(
                {
                    "full_name": "Updated Custodian A",
                    "category": "RELATIVE_OF_EMPLOYEE",
                    "related_employee": "Brother of VP Procurement",
                    "department": "Internal Audit",
                    "designation": "Director",
                    "is_substantiated": True,
                    "status": "FLAGGED",
                    "keywords": ["kickback", "shell company"],
                    "notes": "Updated investigation notes",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["profile"]["full_name"], "Updated Custodian A")
        self.assertEqual(data["profile"]["category"], "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(data["profile"]["related_employee"], "Brother of VP Procurement")
        self.assertEqual(data["profile"]["department"], "Internal Audit")
        self.assertTrue(data["profile"]["is_substantiated"])
        self.assertEqual(data["profile"]["status"], "FLAGGED")
        self.assertEqual(data["profile"]["keywords"], ["kickback", "shell company"])

        # Check DB updated
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.full_name, "Updated Custodian A")
        self.assertEqual(self.profile.category, "RELATIVE_OF_EMPLOYEE")
        self.assertEqual(self.profile.related_employee, "Brother of VP Procurement")
        self.assertTrue(self.profile.is_substantiated)
        self.assertEqual(self.profile.status, "FLAGGED")
        self.assertEqual(self.profile.keywords, ["kickback", "shell company"])

    def test_edit_profile_view_form_post(self):
        url = reverse("edit_profile", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data={
                "full_name": "Form Edited Target",
                "category": "VENDOR",
                "department": "Finance",
                "designation": "CFO",
                "is_substantiated": "true",
                "keywords": json.dumps(["bribe", "hawala"]),
                "next": "/",
            },
        )
        self.assertEqual(res.status_code, 302)
        self.profile.refresh_from_db()
        self.assertEqual(self.profile.full_name, "Form Edited Target")
        self.assertEqual(self.profile.category, "VENDOR")
        self.assertTrue(self.profile.is_substantiated)
        self.assertEqual(self.profile.keywords, ["bribe", "hawala"])

    def test_edit_profile_view_empty_name(self):
        url = reverse("edit_profile", kwargs={"profile_id": str(self.profile.id)})
        res = self.client.post(
            url,
            data=json.dumps({"full_name": "   "}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 400)
        self.assertIn("required", res.json()["message"].lower())

    def test_edit_profile_view_not_found(self):
        url = reverse("edit_profile", kwargs={"profile_id": str(uuid.uuid4())})
        res = self.client.post(
            url,
            data=json.dumps({"full_name": "Ghost Profile"}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 404)

    def test_sync_global_profiles_from_apps(self):
        # Test the sync behavior for q_verify and q_voice
        # Since these apps are installed, the imports should succeed and the loops will run.
        from q_verify.models import VerificationCase

        from .profiles import sync_all_existing_entities_to_profiles

        VerificationCase.objects.create(
            case_ref="TEST-VER-1",
            case_title="Sync Test Case",
            custodian_name="Test Sync Custodian",
            custodian_email="sync@test.com",
            custodian_department="IT",
        )

        from django.utils import timezone
        from q_voice.models import AudioRecording

        AudioRecording.objects.create(
            source_filename="test_sync.wav",
            custodian_name="Test Voice Sync Custodian",
            risk_score=60,
            call_timestamp=timezone.now(),
        )

        from q_bank.models import AuditedPerson

        AuditedPerson.objects.create(
            full_name="Test Bank Sync Custodian",
            notes="flagged",
        )

        created = sync_all_existing_entities_to_profiles()
        self.assertGreaterEqual(created, 3)


class ProfileKeywordRegistryIntegrationTests(TestCase):
    """
    Validates end-to-end profile keyword registry consumption across
    all forensic Q-apps (q_bank, q_mail, q_voice, q_chat, q_trail, q_scan).
    """

    def setUp(self):
        self.profile = InvestigationProfile.objects.create(
            full_name="Vikram Sethi",
            department="Procurement",
            keywords=["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"],
        )
        self.factory = RequestFactory()

    def test_get_profile_keywords_resolver(self):
        from core.profiles import get_profile_keywords

        # 1. By Profile ID
        kws = get_profile_keywords(profile_id=self.profile.id)
        self.assertEqual(
            kws, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

        # 2. By Custodian Name with (Auditee) suffix
        kws_auditee = get_profile_keywords(custodian_name="Vikram Sethi (Auditee)")
        self.assertEqual(
            kws_auditee, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

        # 3. By Request Session Active Profile
        req = self.factory.get("/")
        req.session = {"active_profile_id": str(self.profile.id)}
        kws_session = get_profile_keywords(request=req)
        self.assertEqual(
            kws_session, ["PROJECT_TITAN_SECRET", "SWISS_ACCOUNT", "COMMISSION_CUT", "SHELL_CORP"]
        )

    def test_q_bank_fuzzy_search_profile_keywords(self):
        from q_bank.models import AuditedPerson, BankAccount, BankTransaction

        person = AuditedPerson.objects.create(full_name="Vikram Sethi")
        account = BankAccount.objects.create(
            person=person,
            account_number="9988776655",
            bank_name="HDFC Bank",
        )
        BankTransaction.objects.create(
            account=account,
            txn_date="2026-03-01T10:00:00Z",
            narration="IMPS PAYMENT FOR PROJECT_TITAN_SECRET ROUTING",
            direction=BankTransaction.Direction.DEBIT,
            debit_amount=Decimal("50000.00"),
        )

        client = Client()
        session = client.session
        session["portal_authenticated"] = True
        session["active_profile_id"] = str(self.profile.id)
        session.save()

        # Call fuzzy search without keywords param; should resolve from profile
        res = client.get(f"/bank/api/fuzzy-search/?person_id={person.id}&threshold=70")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertGreaterEqual(data["total_matches"], 1)

    def test_q_mail_checkpoint_evaluation_profile_keywords(self):
        from q_mail.backend.checkpoints import evaluate_email_checkpoints

        class MockEmail:
            subject = "Confidential: Payment and PROJECT_TITAN_SECRET confirmation"
            body_plain = "Please ensure the shell_corp documentation is attached."
            sender_email = "vikram@external.com"
            recipients_to = ["auditor@hmil.net"]
            recipients_cc = []
            recipients_bcc = []

        res = evaluate_email_checkpoints(MockEmail(), profile_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", res["matched_profile_keywords"])
        self.assertIn("SHELL_CORP", res["matched_profile_keywords"])
        badges = [b["label"] for b in res["badges"]]
        self.assertIn("Profile: PROJECT_TITAN_SECRET", badges)
        self.assertIn("Profile: SHELL_CORP", badges)

    def test_q_voice_tagging_and_screening_profile_keywords(self):
        from q_voice.backend.voice_parser import (
            screen_text_for_intent,
            tag_transcript_detections,
        )

        text = "He mentioned using project_titan_secret to dispatch the parcel."
        dets = tag_transcript_detections(text, extra_keywords=self.profile.keywords)
        types = [d["type"] for d in dets]
        self.assertIn("profile", types)
        terms = [d["term"] for d in dets]
        self.assertIn("PROJECT_TITAN_SECRET", terms)

        intent, flagged, risk = screen_text_for_intent(text, extra_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", flagged)
        self.assertGreaterEqual(risk, 35)

    def test_q_chat_screening_profile_keywords(self):
        from q_chat.backend.chat_parser import screen_message_text

        text = "Meet me near the hotel, bring the project_titan_secret voucher slip."
        risk, flagged = screen_message_text(text, extra_keywords=self.profile.keywords)
        self.assertIn("PROJECT_TITAN_SECRET", flagged)
        self.assertGreaterEqual(risk, 25)

    def test_q_trail_keyword_matching(self):
        from q_trail.services import analyze_profiles_money_trail

        profile2 = InvestigationProfile.objects.create(
            full_name="Rajesh Sharma",
            department="Vendor Management",
            keywords=["OFFSHORE_ACC"],
        )

        res = analyze_profiles_money_trail(profile_ids=[str(self.profile.id), str(profile2.id)])
        self.assertIn("trail_keywords", res)
        self.assertIn("PROJECT_TITAN_SECRET", res["trail_keywords"])
        self.assertIn("OFFSHORE_ACC", res["trail_keywords"])

    def test_q_scan_config_profile_keywords_injection(self):
        client = Client()
        session = client.session
        session["portal_authenticated"] = True
        session["active_profile_id"] = str(self.profile.id)
        session.save()

        res = client.get("/scan/download/config.json/")
        self.assertEqual(res.status_code, 200)
        cfg_data = json.loads(res.content.decode("utf-8"))
        self.assertIn("PROJECT_TITAN_SECRET", cfg_data["keywords"])
        self.assertIn("SWISS_ACCOUNT", cfg_data["keywords"])


class CoreFuzzyEngineTests(TestCase):
    def test_extract_keywords_from_string(self):
        from core.fuzzy import extract_keywords_from_string

        self.assertEqual(extract_keywords_from_string(""), [])
        res = extract_keywords_from_string("alpha, beta; gamma\r\nALPHA, delta")
        self.assertEqual(res, ["alpha", "beta", "gamma", "delta"])

    def test_extract_keywords_from_txt(self):
        from core.fuzzy import extract_keywords_from_txt

        raw_bytes = b"# Comment line\n// Another comment\napple, banana\n\ncherry"
        res = extract_keywords_from_txt(raw_bytes)
        self.assertEqual(res, ["apple", "banana", "cherry"])

        # String input
        res_str = extract_keywords_from_txt("dog, cat\n# ignore\nfish")
        self.assertEqual(res_str, ["dog", "cat", "fish"])

        # Latin-1 encoded bytes
        latin1_bytes = "café, naïve".encode("latin-1")
        res_latin = extract_keywords_from_txt(latin1_bytes)
        self.assertIn("café", res_latin)

    def test_extract_keywords_from_xlsx_general_table(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header with stopword
        ws.append(["keyword", "notes"])
        ws.append(["shell_corp", "internal alert"])
        ws.append(["bribe_fund", "cash"])
        ws.append([None, None])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("shell_corp", res)
        self.assertIn("bribe_fund", res)

    def test_extract_keywords_from_xlsx_targeted_header(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Target Name", "Watchlist Identifier", "Date"])
        ws.append(["Vikram", "ACC-9988; HAWALA-1", "2026-01-01"])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("ACC-9988", res)
        self.assertIn("HAWALA-1", res)

    def test_extract_keywords_from_xlsx_fallback_no_target_col(self):
        import io

        import openpyxl

        from core.fuzzy import extract_keywords_from_xlsx

        wb = openpyxl.Workbook()
        ws = wb.active
        # Header without any target column words
        ws.append(["Category", "Details"])
        ws.append(["Finance", "hawala_node, shell_account"])
        buf = io.BytesIO()
        wb.save(buf)

        res = extract_keywords_from_xlsx(buf.getvalue())
        self.assertIn("hawala_node", res)
        self.assertIn("shell_account", res)

    def test_extract_keywords_from_file_and_fallbacks(self):
        from core.fuzzy import extract_keywords_from_file

        # TXT file obj
        txt_file = SimpleUploadedFile("watch.csv", b"wire_transfer, kickback")
        res_txt = extract_keywords_from_file(txt_file, "watch.csv")
        self.assertIn("wire_transfer", res_txt)

        # Fallback unknown extension
        unknown_file = SimpleUploadedFile("unknown.dat", b"secret_fund, offshore")
        res_unknown = extract_keywords_from_file(unknown_file, "unknown.dat")
        self.assertIn("secret_fund", res_unknown)

    def test_extract_keywords_from_request(self):
        from core.fuzzy import extract_keywords_from_request

        factory = RequestFactory()

        # 1. From GET query
        req_get = factory.get("/?keywords=alpha,beta")
        res_get = extract_keywords_from_request(req_get)
        self.assertEqual(res_get, ["alpha", "beta"])

        # 2. From POST q
        req_post = factory.post("/", {"q": "gamma; delta"})
        res_post = extract_keywords_from_request(req_post)
        self.assertEqual(res_post, ["gamma", "delta"])

        # 3. From attached file
        f = SimpleUploadedFile("terms.txt", b"epsilon, zeta")
        req_file = factory.post("/", {"file": f})
        res_file = extract_keywords_from_request(req_file)
        self.assertIn("epsilon", res_file)
        self.assertIn("zeta", res_file)

    def test_score_text_against_keywords(self):
        from core.fuzzy import score_text_against_keywords

        # Empty cases
        self.assertEqual(score_text_against_keywords("", ["abc"]), (False, 0, ""))
        self.assertEqual(score_text_against_keywords("some text", []), (False, 0, ""))

        # Direct 100% substring match
        matched, score, kw = score_text_against_keywords(
            "Payment to Apex Logistics", ["Apex Logistics"]
        )
        self.assertTrue(matched)
        self.assertEqual(score, 100)
        self.assertEqual(kw, "Apex Logistics")

        # Fuzzy token match
        matched_fuzz, score_fuzz, kw_fuzz = score_text_against_keywords(
            "Transfer to Logistix", ["Logistics"], threshold=70
        )
        self.assertTrue(matched_fuzz)
        self.assertGreaterEqual(score_fuzz, 70)
        self.assertEqual(kw_fuzz, "Logistics")

        # Below threshold match
        matched_no, _, _ = score_text_against_keywords(
            "Completely unrelated text", ["Logistics"], threshold=90
        )
        self.assertFalse(matched_no)


class CoreAuditTests(TestCase):
    def setUp(self):
        from core.models import Audit, InvestigationProfile

        self.Audit = Audit
        self.InvestigationProfile = InvestigationProfile
        self.client = Client()
        session = self.client.session
        session["portal_authenticated"] = True
        session.save()

        # Create sample profiles
        self.profile1 = self.InvestigationProfile.objects.create(
            full_name="Arun Kumar",
            employee_id="EMP-1001",
            department="Procurement",
            designation="Manager",
            status="ACTIVE",
            is_substantiated=True,
        )
        self.profile2 = self.InvestigationProfile.objects.create(
            full_name="Rajesh Sharma",
            employee_id="EMP-1002",
            department="Logistics",
            designation="Executive",
            status="ACTIVE",
            is_substantiated=False,
        )
        self.profile3 = self.InvestigationProfile.objects.create(
            full_name="Priya Patel",
            employee_id="EMP-1003",
            department="Finance",
            designation="Analyst",
            status="ACTIVE",
            is_substantiated=False,
        )

    def test_auto_generate_audit_name_sequence(self):
        from core.audits import create_audit, generate_next_audit_name

        # First audit in 2026 starts at 2026-WB-01
        name1 = generate_next_audit_name(year=2026)
        self.assertEqual(name1, "2026-WB-01")

        audit1 = create_audit(year=2026, title="Audit 1")
        self.assertEqual(audit1.name, "2026-WB-01")

        # Second audit increments to 2026-WB-02
        name2 = generate_next_audit_name(year=2026)
        self.assertEqual(name2, "2026-WB-02")

        audit2 = create_audit(year=2026, title="Audit 2")
        self.assertEqual(audit2.name, "2026-WB-02")

        # Third audit increments to 2026-WB-03
        audit3 = create_audit(year=2026, title="Audit 3")
        self.assertEqual(audit3.name, "2026-WB-03")

    def test_audit_name_generation_across_years(self):
        from core.audits import create_audit, generate_next_audit_name

        create_audit(year=2025, title="2025 Audit")
        self.assertEqual(generate_next_audit_name(year=2025), "2025-WB-02")
        # 2026 starts independently at 01
        self.assertEqual(generate_next_audit_name(year=2026), "2026-WB-01")

    def test_audit_name_generation_expansion_beyond_99(self):
        from core.audits import generate_next_audit_name

        self.Audit.objects.create(name="2026-WB-99", title="High Seq Audit")
        self.assertEqual(generate_next_audit_name(year=2026), "2026-WB-100")

    def test_create_audit_with_mapped_profiles(self):
        from core.audits import create_audit

        audit = create_audit(
            title="Procurement Collusion Review",
            description="Investigating supplier kickbacks and fake invoices.",
            status="ACTIVE",
            profile_ids=[str(self.profile1.id), str(self.profile2.id)],
        )

        self.assertEqual(audit.profiles.count(), 2)
        self.assertIn(self.profile1, audit.profiles.all())
        self.assertIn(self.profile2, audit.profiles.all())
        self.assertNotIn(self.profile3, audit.profiles.all())

        # Check serialization in to_dict
        d = audit.to_dict()
        self.assertEqual(d["name"], audit.name)
        self.assertEqual(d["title"], "Procurement Collusion Review")
        self.assertEqual(d["status"], "ACTIVE")
        self.assertEqual(d["profiles_count"], 2)
        self.assertIn(str(self.profile1.id), d["profile_ids"])

        # Check profile to_dict audits reference
        p1_dict = self.profile1.to_dict()
        self.assertEqual(len(p1_dict["audits"]), 1)
        self.assertEqual(p1_dict["audits"][0]["name"], audit.name)

    def test_map_and_unmap_profiles(self):
        from core.audits import create_audit, map_profiles_to_audit, unmap_profile_from_audit

        audit = create_audit(title="Logistics Review", profile_ids=[str(self.profile1.id)])
        self.assertEqual(audit.profiles.count(), 1)

        # Add profile 2
        map_profiles_to_audit(audit.id, [str(self.profile2.id)], replace=False)
        audit.refresh_from_db()
        self.assertEqual(audit.profiles.count(), 2)

        # Replace with only profile 3
        map_profiles_to_audit(audit.id, [str(self.profile3.id)], replace=True)
        audit.refresh_from_db()
        self.assertEqual(audit.profiles.count(), 1)
        self.assertEqual(audit.profiles.first().id, self.profile3.id)

        # Unmap profile 3
        unmap_profile_from_audit(audit.id, self.profile3.id)
        audit.refresh_from_db()
        self.assertEqual(audit.profiles.count(), 0)

    def test_landing_view_contains_audits(self):
        from core.audits import create_audit

        create_audit(title="Landing Test Audit", profile_ids=[str(self.profile1.id)])

        response = self.client.get(reverse("landing"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("audits", response.context)
        self.assertIn("audits_json", response.context)
        self.assertIn("next_audit_name", response.context)
        self.assertGreaterEqual(len(response.context["audits"]), 1)

    def test_create_audit_view_json(self):
        payload = {
            "title": "API Created Audit",
            "description": "Via JSON endpoint",
            "status": "IN_PROGRESS",
            "profile_ids": [str(self.profile1.id), str(self.profile2.id)],
        }
        response = self.client.post(
            reverse("create_audit"),
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["audit"]["title"], "API Created Audit")
        self.assertEqual(data["audit"]["status"], "IN_PROGRESS")
        self.assertEqual(data["audit"]["profiles_count"], 2)
        self.assertTrue(data["audit"]["name"].endswith("-01") or "-WB-" in data["audit"]["name"])

    def test_create_audit_view_json_with_null_name(self):
        payload = {
            "name": None,
            "title": "Null Name Audit",
            "description": None,
            "status": None,
            "profile_ids": [str(self.profile1.id)],
        }
        response = self.client.post(
            reverse("create_audit"),
            data=json.dumps(payload),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["audit"]["title"], "Null Name Audit")
        self.assertEqual(data["audit"]["status"], "ACTIVE")
        self.assertTrue("-WB-" in data["audit"]["name"])

    def test_create_audit_view_form_post(self):
        # 1. Standard list in POST
        response = self.client.post(
            reverse("create_audit"),
            data={
                "title": "Form Created Audit",
                "status": "ACTIVE",
                "profile_ids": [str(self.profile1.id)],
                "next": "/",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(self.Audit.objects.filter(title="Form Created Audit").exists())

        # 2. Comma-separated string in POST
        response_comma = self.client.post(
            reverse("create_audit"),
            data={
                "title": "Comma Created Audit",
                "status": "ACTIVE",
                "profile_ids": f"{self.profile1.id},{self.profile2.id}",
                "next": "/",
            },
        )
        self.assertEqual(response_comma.status_code, 302)
        self.assertTrue(self.Audit.objects.filter(title="Comma Created Audit").exists())

    def test_map_audit_profiles_view(self):
        from core.audits import create_audit

        audit = create_audit(title="Mapping View Test")
        response = self.client.post(
            reverse("map_audit_profiles", kwargs={"audit_id": audit.id}),
            data=json.dumps({"profile_ids": [str(self.profile1.id), str(self.profile3.id)]}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["audit"]["profiles_count"], 2)

        # String json profile_ids
        res_str = self.client.post(
            reverse("map_audit_profiles", kwargs={"audit_id": audit.id}),
            data=json.dumps({"profile_ids": f'["{self.profile1.id}"]'}),
            content_type="application/json",
        )
        self.assertEqual(res_str.status_code, 200)

        # 404 for non-existent audit
        res_404 = self.client.post(
            reverse("map_audit_profiles", kwargs={"audit_id": uuid.uuid4()}),
            data=json.dumps({"profile_ids": []}),
            content_type="application/json",
        )
        self.assertEqual(res_404.status_code, 404)

    def test_api_audits_list_and_next_name(self):
        from core.audits import create_audit

        create_audit(title="List Test Audit")

        # List endpoint
        res_list = self.client.get(reverse("api_audits"))
        self.assertEqual(res_list.status_code, 200)
        data_list = res_list.json()
        self.assertEqual(data_list["status"], "success")
        self.assertGreaterEqual(len(data_list["audits"]), 1)

        # Next name endpoint
        res_name = self.client.get(reverse("api_next_audit_name"))
        self.assertEqual(res_name.status_code, 200)
        data_name = res_name.json()
        self.assertEqual(data_name["status"], "success")
        self.assertTrue("-WB-" in data_name["next_name"])

    def test_set_active_audit_and_context_filtering(self):
        from django.test import RequestFactory

        from core.audits import create_audit
        from core.context_processors import global_profiles_context

        audit = create_audit(
            title="Active Filter Test Audit",
            profile_ids=[str(self.profile1.id), str(self.profile2.id)],
        )

        # 1. API set active audit via POST
        res = self.client.post(
            reverse("set_active_audit"),
            data=json.dumps({"audit_id": str(audit.id)}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["active_audit"]["name"], audit.name)

        # 2. Context processor reflects filtered profiles
        factory = RequestFactory()
        req = factory.get("/")
        req.session = self.client.session
        ctx = global_profiles_context(req)

        self.assertTrue(ctx["is_audit_active"])
        self.assertEqual(ctx["active_audit"].id, audit.id)
        # Should only have profile1 and profile2, not profile3
        self.assertEqual(len(ctx["investigation_profiles"]), 2)
        self.assertEqual(ctx["audit_profiles_count"], 2)
        self.assertGreaterEqual(ctx["total_profiles_count"], 3)
        prof_ids = [p.id for p in ctx["investigation_profiles"]]
        self.assertIn(self.profile1.id, prof_ids)
        self.assertIn(self.profile2.id, prof_ids)
        self.assertNotIn(self.profile3.id, prof_ids)

        # 3. Clear active audit
        res_clear = self.client.post(
            reverse("set_active_audit"),
            data=json.dumps({"audit_id": ""}),
            content_type="application/json",
        )
        self.assertEqual(res_clear.status_code, 200)
        clear_data = res_clear.json()
        self.assertEqual(clear_data["status"], "success")
        self.assertIsNone(clear_data["active_audit"])

        # 4. Form POST set active audit with redirect
        res_form = self.client.post(
            reverse("set_active_audit"),
            data={"audit_id": str(audit.id), "next": "/"},
        )
        self.assertEqual(res_form.status_code, 302)
        self.assertEqual(res_form.url, "/")

        req.session = self.client.session
        ctx_cleared = global_profiles_context(req)
        self.assertTrue(ctx_cleared["is_audit_active"])
        self.assertEqual(ctx_cleared["active_audit"].id, audit.id)

    def test_api_profiles_list(self):
        res = self.client.get(reverse("api_profiles"))
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertGreaterEqual(len(data["profiles"]), 3)

    def test_set_active_profile_view_json_and_redirect(self):
        # 1. JSON POST to set active profile
        res_json = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": str(self.profile1.id)}),
            content_type="application/json",
        )
        self.assertEqual(res_json.status_code, 200)
        data = res_json.json()
        self.assertEqual(data["status"], "success")
        self.assertEqual(data["active_profile"]["full_name"], "Arun Kumar")

        # 2. JSON POST with empty profile_id to clear
        res_clear = self.client.post(
            reverse("set_active_profile"),
            data=json.dumps({"profile_id": ""}),
            content_type="application/json",
        )
        self.assertEqual(res_clear.status_code, 200)
        self.assertIsNone(res_clear.json()["active_profile"])

        # 3. Form POST to set and redirect
        res_form = self.client.post(
            reverse("set_active_profile"),
            data={"profile_id": str(self.profile2.id), "next": "/"},
        )
        self.assertEqual(res_form.status_code, 302)
        self.assertEqual(res_form.url, "/")

    def test_append_profile_keywords_view(self):
        url = reverse("add_profile_keywords", kwargs={"profile_id": self.profile1.id})
        res = self.client.post(
            url,
            data=json.dumps({"keywords": ["shell vendor", "kickback"]}),
            content_type="application/json",
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("shell vendor", data["profile"]["keywords"])

        # Form POST submission
        res_post = self.client.post(
            url,
            data={"keywords": "bribe, siphoning"},
        )
        self.assertEqual(res_post.status_code, 200)

        # Error: missing keywords
        res_bad = self.client.post(
            url,
            data=json.dumps({"keywords": []}),
            content_type="application/json",
        )
        self.assertEqual(res_bad.status_code, 400)

        # Error: non-existent profile
        bad_url = reverse("add_profile_keywords", kwargs={"profile_id": uuid.uuid4()})
        res_404 = self.client.post(
            bad_url,
            data=json.dumps({"keywords": ["test"]}),
            content_type="application/json",
        )
        self.assertEqual(res_404.status_code, 404)

    def test_parse_keywords_file_view(self):
        url = reverse("parse_keywords_file")

        # 1. Plain text file
        txt_content = b'"offshore account"\nhawala\nbribe\n'
        txt_file = SimpleUploadedFile("watchlist.txt", txt_content, content_type="text/plain")
        res_txt = self.client.post(url, {"file": txt_file})
        self.assertEqual(res_txt.status_code, 200)
        data_txt = res_txt.json()
        self.assertEqual(data_txt["status"], "success")
        self.assertIn("offshore account", data_txt["keywords"])

        # 2. CSV file (rejected because only .txt is allowed)
        csv_content = b"keyword\nfront company\nsiphoning\n"
        csv_file = SimpleUploadedFile("watchlist.csv", csv_content, content_type="text/csv")
        res_csv = self.client.post(url, {"file": csv_file})
        self.assertEqual(res_csv.status_code, 400)
        self.assertIn("Only .txt files are supported", res_csv.json()["message"])

        # 3. Missing file
        res_missing = self.client.post(url, {})
        self.assertEqual(res_missing.status_code, 400)

        # 4. Invalid file extension
        pdf_file = SimpleUploadedFile("watchlist.pdf", b"%PDF-1.4", content_type="application/pdf")
        res_invalid = self.client.post(url, {"file": pdf_file})
        self.assertEqual(res_invalid.status_code, 400)

    def test_upload_profile_keywords_file_view(self):
        url = reverse("upload_profile_keywords_file", kwargs={"profile_id": self.profile1.id})

        txt_file = SimpleUploadedFile(
            "keywords.txt", b"bribe\nkickback\n", content_type="text/plain"
        )
        res = self.client.post(url, {"file": txt_file})
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["status"], "success")
        self.assertIn("kickback", data["profile"]["keywords"])

        # Missing file
        res_missing = self.client.post(url, {})
        self.assertEqual(res_missing.status_code, 400)

        # Unsupported extension
        res_unsupp = self.client.post(
            url,
            {
                "file": SimpleUploadedFile(
                    "bad.exe", b"exe", content_type="application/octet-stream"
                )
            },
        )
        self.assertEqual(res_unsupp.status_code, 400)

        # Empty keywords in file
        res_empty = self.client.post(
            url,
            {"file": SimpleUploadedFile("empty.txt", b"\n\n", content_type="text/plain")},
        )
        self.assertEqual(res_empty.status_code, 400)

        # Non-existent profile
        bad_url = reverse("upload_profile_keywords_file", kwargs={"profile_id": uuid.uuid4()})
        res_404 = self.client.post(
            bad_url,
            {"file": SimpleUploadedFile("k.txt", b"kw", content_type="text/plain")},
        )
        self.assertEqual(res_404.status_code, 404)

    def test_core_profiles_extract_keywords_from_excel_and_helpers(self):
        import openpyxl

        from core.profiles import extract_keywords_from_file

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Watchlist"
        ws.append(["Keyword", "Risk Level"])
        ws.append(["shell company", "HIGH"])
        ws.append(["fake invoice, kickback", "CRITICAL"])
        ws.append(["", ""])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        kws = extract_keywords_from_file(buf, "watchlist.xlsx")
        self.assertIn("shell company", kws)
        self.assertIn("fake invoice", kws)
        self.assertIn("kickback", kws)

    def test_core_profiles_sync_profiles_from_all_modules(self):
        from q_bank.models import AuditedPerson
        from q_verify.models import VerificationCase
        from q_voice.models import AudioRecording

        from core.profiles import sync_all_existing_entities_to_profiles

        # Q-Bank person
        AuditedPerson.objects.create(
            full_name="Vikram Seth (Auditee)",
            employee_id="EMP-9901",
            department="Procurement",
            designation="Director",
            notes="Flagged for audit review",
        )
        # Q-Voice recording
        from django.utils import timezone

        AudioRecording.objects.create(
            call_ref="CALL-9901",
            call_title="Suspicious wiretap call",
            call_timestamp=timezone.now(),
            duration_seconds=120,
            custodian_name="Sanjay Verma",
            risk_score=75,
        )
        # Q-Verify case
        VerificationCase.objects.create(
            case_ref="VER-9901",
            case_title="Supplier Invoices Audit",
            custodian_name="Anil Kapoor",
            custodian_department="Supply Chain",
            custodian_email="anil@example.com",
        )

        created = sync_all_existing_entities_to_profiles()
        self.assertGreaterEqual(created, 3)

        self.assertTrue(self.InvestigationProfile.objects.filter(full_name="Vikram Seth").exists())
        self.assertTrue(self.InvestigationProfile.objects.filter(full_name="Sanjay Verma").exists())
        self.assertTrue(self.InvestigationProfile.objects.filter(full_name="Anil Kapoor").exists())

    def test_resolve_or_create_profile_from_request(self):
        from django.contrib.sessions.backends.db import SessionStore

        from core.profiles import resolve_or_create_profile_from_request

        factory = RequestFactory()

        # 1. Existing profile by ID
        req1 = factory.post("/", {"profile_id": str(self.profile1.id)})
        req1.session = SessionStore()
        p1, name1 = resolve_or_create_profile_from_request(req1)
        self.assertEqual(p1.id, self.profile1.id)
        self.assertEqual(name1, self.profile1.full_name)

        # 2. Inline new profile
        req2 = factory.post(
            "/",
            {
                "new_profile_name": "New Investigator Subject",
                "new_profile_dept": "Internal Audit",
                "new_profile_role": "Specialist",
            },
        )
        req2.session = SessionStore()
        p2, name2 = resolve_or_create_profile_from_request(req2)
        self.assertEqual(name2, "New Investigator Subject")
        self.assertEqual(p2.department, "Internal Audit")

        # 3. Fallback custodian name
        req3 = factory.post("/", {"custodian_name": "Fallback Custodian"})
        req3.session = SessionStore()
        p3, name3 = resolve_or_create_profile_from_request(req3)
        self.assertIsNotNone(p3)
        self.assertEqual(name3, "Fallback Custodian")


class CorePromptsEngineTests(TestCase):
    """Verifies decoupled prompt loading, rendering, caching, and fallback handling."""

    def setUp(self):
        from core.prompts import clear_prompt_cache

        clear_prompt_cache()

    def test_load_q_trail_prompts(self):
        from core.prompts import load_prompt

        sys_prompt = load_prompt("q_trail", "loop_narrative_system.txt")
        self.assertIn("forensic financial intelligence auditor", sys_prompt)

        user_prompt = load_prompt("q_trail", "loop_narrative_user.txt")
        self.assertIn("Closed circular round-tripping loop", user_prompt)
        self.assertIn("{cycle_str}", user_prompt)

    def test_load_q_link_prompts(self):
        from core.prompts import load_prompt

        single_prompt = load_prompt("q_link", "single_entity_dossier.txt")
        self.assertIn("ForensiQ Copilot", single_prompt)
        self.assertIn("{target_name}", single_prompt)

        dual_prompt = load_prompt("q_link", "dual_entity_pathways.txt")
        self.assertIn("{source_name}", dual_prompt)
        self.assertIn("{target_name}", dual_prompt)

        overview_prompt = load_prompt("q_link", "syndicate_overview.txt")
        self.assertIn("{total_entities}", overview_prompt)

    def test_render_prompt_with_context(self):
        from core.prompts import render_prompt

        rendered = render_prompt(
            "q_trail",
            "loop_narrative_user.txt",
            cycle_str="A -> B -> A",
            initial_amt="10,000.00",
            return_amt="9,500.00",
            retained_amt="500.00",
            conduits_str="Bank B",
        )
        self.assertIn("Path: A -> B -> A", rendered)
        self.assertIn("Initial Dispatched Outflow: ₹10,000.00", rendered)
        self.assertIn("Return Inflow: ₹9,500.00", rendered)
        self.assertIn("Conduits Withheld Fee: ₹500.00", rendered)

    def test_missing_prompt_uses_fallback(self):
        from core.prompts import load_prompt, render_prompt

        fallback = "Custom fallback prompt content."
        loaded = load_prompt("unknown_app", "nonexistent.txt", fallback=fallback)
        self.assertEqual(loaded, fallback)

        rendered = render_prompt("unknown_app", "nonexistent.txt", fallback=fallback)
        self.assertEqual(rendered, fallback)

    def test_render_missing_keys_safe_behavior(self):
        from core.prompts import render_prompt

        # Missing one placeholder should not crash, leaves other placeholders formatted or safe
        rendered = render_prompt(
            "q_trail",
            "loop_narrative_user.txt",
            cycle_str="Node1 -> Node2",
            # deliberately omitting other placeholders
        )
        self.assertIn("Node1 -> Node2", rendered)

    def test_prompt_path_resolves_to_apps_dir(self):
        from core.prompts import get_prompt_path

        p = get_prompt_path("q_trail", "loop_narrative_system.txt")
        self.assertTrue(p.exists())
        self.assertIn("apps", str(p).lower())
        self.assertIn("q_trail", str(p).lower())
        self.assertIn("prompts", str(p).lower())
